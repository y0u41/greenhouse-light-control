#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
md2docx.py —— 由 Markdown 生成规范格式的项目文档 DOCX

用途
    以 <项目根>/项目规划书.md 为唯一事实来源，生成同名的 .docx 提交稿。
    改文档只改 .md，然后重跑本脚本，避免 md 与 docx 两份内容漂移。

用法
    py tools/md2docx.py
    py tools/md2docx.py --input 项目规划书.md --output 项目规划书.docx

格式约定（依据 项目规划书.md §15.2）
    · A4，页边距 2 cm
    · 正文：宋体 / Times New Roman，小四（12 pt），行距 1.5 倍
    · 标题：黑体，H1 18 pt / H2 15 pt / H3 13 pt
    · 代码：Consolas 五号（10.5 pt），灰底 + 细边框，单倍行距
    · 表格：Table Grid 边框，首行加粗 + 底纹 + 跨页重复
    · 页脚居中页码

支持的 Markdown 子集
    标题 #/##/###、段落、围栏代码块、表格、引用块、无序列表、有序列表
    （有序列表保留原文数字，交给 Word 自动编号会与手写编号重复）
    任务列表 - [ ] / - [x]、分隔线 ---、行内 **加粗** 与 `代码`

依赖
    python-docx（随 DSH 运行时提供，无需另行安装）
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:  # pragma: no cover
    sys.exit("缺少 python-docx，请使用 DSH 运行时提供的 Python 解释器运行本脚本。")


# ---------------------------------------------------------------- 样式常量
FONT_EA_BODY = "宋体"
FONT_EA_HEAD = "黑体"
FONT_ASCII_BODY = "Times New Roman"
FONT_CODE = "Consolas"

SIZE_BODY = 12.0        # 小四
SIZE_H1 = 18.0
SIZE_H2 = 15.0
SIZE_H3 = 13.0
SIZE_H4 = 12.0          # 四级标题：与正文同号但加粗着色，保持层级可辨
SIZE_CODE = 10.5        # 五号
SIZE_TABLE = 10.5

COLOR_HEAD = "1F3864"   # 深蓝，正式文档常用
COLOR_H4 = "2E5496"
COLOR_MUTED = "595959"
FILL_CODE = "F2F2F2"
FILL_THEAD = "DCE6F5"
COLOR_RULE = "BFBFBF"

INLINE_RE = re.compile(r"(\[[^\]]+\]\([^)]+\)|\*\*.+?\*\*|~~.+?~~|`[^`]+`)")
LINK_RE = re.compile(r"^\[([^\]]+)\]\(([^)]+)\)$")
COLOR_LINK = "1F4E79"


# ---------------------------------------------------------------- 底层工具
def set_run_font(run, ascii_font, ea_font, size=None, bold=None,
                 italic=None, color=None):
    """同时设置西文与东亚字体（只设 run.font.name 对中文无效）。"""
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), ea_font)
    rfonts.set(qn("w:cs"), ascii_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def set_paragraph_borders(paragraph, edges=("bottom",), sz=6, color=COLOR_RULE):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    for edge in edges:
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:space"), "1")
        el.set(qn("w:color"), color)
        pbdr.append(el)


def shade_cell(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def repeat_header_row(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trpr.append(el)


def prevent_row_split(row):
    """禁止单元格内容被分页截断（否则一行会横跨两页）。"""
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:cantSplit"))


def set_table_full_width(table):
    tblpr = table._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:type"), "pct")
    tblw.set(qn("w:w"), "5000")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "autofit")
    tblpr.append(layout)


def add_hyperlink(paragraph, url, text, size=SIZE_BODY, ea=FONT_EA_BODY,
                  color=COLOR_LINK):
    """插入真实超链接（仅用于 http(s)/mailto）。"""
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_ASCII_BODY)
    rfonts.set(qn("w:hAnsi"), FONT_ASCII_BODY)
    rfonts.set(qn("w:eastAsia"), ea)
    rpr.append(rfonts)
    col = OxmlElement("w:color")
    col.set(qn("w:val"), color)
    rpr.append(col)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def set_core_properties(doc, title, subtitle="", source_name=""):
    """写入文档属性，避免提交稿里残留 'python-docx' 之类的生成器署名。"""
    cp = doc.core_properties
    cp.title = title if not subtitle else f"{title} —— {subtitle}"
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ("由 tools/md2docx.py 从 %s 生成，请勿手工编辑本文件；"
                   "改内容请改 .md 后重新生成" % (source_name or "Markdown"))
    cp.revision = 1


def add_page_number_footer(section):
    para = section.footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    set_run_font(run, FONT_ASCII_BODY, FONT_EA_BODY, size=9, color=COLOR_MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, end):
        run._r.append(el)


# ---------------------------------------------------------------- 文档骨架
def build_document():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)          # A4
    section.page_height = Cm(29.7)
    for attr in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.0))

    normal = doc.styles["Normal"]
    normal.font.name = FONT_ASCII_BODY
    normal.font.size = Pt(SIZE_BODY)
    rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), FONT_ASCII_BODY)
    rfonts.set(qn("w:hAnsi"), FONT_ASCII_BODY)
    rfonts.set(qn("w:eastAsia"), FONT_EA_BODY)
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(4)

    for name, size in (("Heading 1", SIZE_H1),
                       ("Heading 2", SIZE_H2),
                       ("Heading 3", SIZE_H3)):
        st = doc.styles[name]
        st.font.name = FONT_EA_HEAD
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(COLOR_HEAD)
        rf = st.element.get_or_add_rPr().get_or_add_rFonts()
        rf.set(qn("w:ascii"), FONT_EA_HEAD)
        rf.set(qn("w:hAnsi"), FONT_EA_HEAD)
        rf.set(qn("w:eastAsia"), FONT_EA_HEAD)

    add_page_number_footer(section)
    normalize_styles(doc)
    return doc


def normalize_styles(doc):
    """把模板里遗留的、本机不存在的字体引用统一掉，避免缺失字体告警。"""
    latent = {"Courier": FONT_CODE, "Courier New": FONT_CODE,
              "Calibri": FONT_ASCII_BODY, "Cambria": FONT_EA_BODY}
    for style in doc.styles:
        el = style.element
        if el is None:
            continue
        for rfonts in el.iter(qn("w:rFonts")):
            for attr in ("w:ascii", "w:hAnsi", "w:cs"):
                cur = rfonts.get(qn(attr))
                if cur in latent:
                    rfonts.set(qn(attr), latent[cur])
            ea = rfonts.get(qn("w:eastAsia"))
            if ea in latent:
                rfonts.set(qn("w:eastAsia"), latent[ea])

    # docDefaults 也指向显式字体，避免任何未显式指定字体的文本落到主题字体上
    styles_el = doc.styles.element
    for rpr_default in styles_el.iter(qn("w:rPrDefault")):
        for rfonts in rpr_default.iter(qn("w:rFonts")):
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                if rfonts.get(qn(attr)) is not None:
                    del rfonts.attrib[qn(attr)]
            rfonts.set(qn("w:ascii"), FONT_ASCII_BODY)
            rfonts.set(qn("w:hAnsi"), FONT_ASCII_BODY)
            rfonts.set(qn("w:eastAsia"), FONT_EA_BODY)
            rfonts.set(qn("w:cs"), FONT_ASCII_BODY)


def add_inline(paragraph, text, size=SIZE_BODY, ea=FONT_EA_BODY,
               ascii_font=FONT_ASCII_BODY, bold=False, italic=False,
               color=None):
    """写入一行文本，处理 **加粗** 与 `行内代码`。"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        link = LINK_RE.match(part)
        if link:
            label, url = link.group(1), link.group(2)
            if url.startswith(("http://", "https://", "mailto:")):
                add_hyperlink(paragraph, url, label, size, ea)
            else:
                # 相对路径的交叉引用（如 [Tech_Design.md](Tech_Design.md)）：
                # 提交稿里只保留可读的引用名，避免出现原始链接语法
                run = paragraph.add_run(label)
                set_run_font(run, ascii_font, ea, size, bold, italic, color)
            continue
        if part.startswith("~~") and part.endswith("~~") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, ascii_font, ea, size, bold, italic, color)
            run.font.strike = True
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            # 递归处理，支持 **粗体里嵌 `代码` 或 [链接](x)**
            add_inline(paragraph, part[2:-2], size, ea, ascii_font,
                       True, italic, color)
            continue
        if part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, FONT_CODE, ea, size - 1.0, bold, italic, color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, ascii_font, ea, size, bold, italic, color)
    return paragraph


# ---------------------------------------------------------------- 解析
def parse_blocks(lines):
    blocks = []
    i, n = 0, len(lines)
    while i < n:
        raw = lines[i].rstrip("\n")
        stripped = raw.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].lstrip().startswith("```"):
                buf.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            blocks.append(("code", buf))
            continue

        if stripped.startswith("|"):
            buf = []
            while i < n and lines[i].lstrip().startswith("|"):
                buf.append(lines[i].strip())
                i += 1
            blocks.append(("table", buf))
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            blocks.append(("h", (len(m.group(1)), m.group(2).strip())))
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", stripped):
            blocks.append(("hr", None))
            i += 1
            continue

        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].lstrip().startswith(">"):
                buf.append(re.sub(r"^\s*>\s?", "", lines[i]).rstrip("\n"))
                i += 1
            blocks.append(("quote", buf))
            continue

        m = re.match(r"^\s*[-*]\s+\[([ xX])\]\s+(.*)$", raw)
        if m:
            blocks.append(("task", (m.group(1).lower() == "x", m.group(2).strip())))
            i += 1
            continue

        m = re.match(r"^\s*[-*]\s+(.*)$", raw)
        if m:
            blocks.append(("ul", m.group(1).strip()))
            i += 1
            continue

        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", raw)
        if m:
            blocks.append(("ol", (m.group(1), m.group(2).strip())))
            i += 1
            continue

        buf = [stripped]
        i += 1
        while i < n:
            nxt = lines[i]
            ns = nxt.strip()
            if (not ns
                    or ns.startswith(("|", ">", "```", "#"))
                    or re.match(r"^\s*[-*]\s+", nxt)
                    or re.match(r"^\s*\d+[.)]\s+", nxt)
                    or re.match(r"^\s*(-{3,}|\*{3,}|_{3,})\s*$", nxt)):
                break
            buf.append(ns)
            i += 1
        blocks.append(("p", " ".join(buf)))

    return blocks


def parse_table(rows):
    grid = []
    for row in rows:
        cells = [c.strip() for c in row.strip().strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
            continue                      # 分隔行
        grid.append(cells)
    if not grid:
        return grid
    width = max(len(r) for r in grid)
    return [r + [""] * (width - len(r)) for r in grid]


# ---------------------------------------------------------------- 渲染
def render(doc, blocks):
    for kind, payload in blocks:

        if kind == "h":
            level, text = payload
            lvl = min(level, 4)
            size = {1: SIZE_H1, 2: SIZE_H2, 3: SIZE_H3}.get(lvl, SIZE_H4)
            para = doc.add_paragraph(style=f"Heading {min(level, 3)}")
            add_inline(para, text, size=size,
                       ea=FONT_EA_HEAD, ascii_font=FONT_EA_HEAD, bold=True,
                       color=COLOR_HEAD if lvl < 4 else COLOR_H4)
            para.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            para.paragraph_format.space_after = Pt(6)
            para.paragraph_format.line_spacing = 1.2
            continue

        if kind == "p":
            para = doc.add_paragraph()
            add_inline(para, payload)
            continue

        if kind == "quote":
            for j, line in enumerate(payload):
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Cm(0.6)
                para.paragraph_format.line_spacing = 1.3
                para.paragraph_format.space_after = Pt(2 if j < len(payload) - 1 else 8)
                if line.strip():
                    add_inline(para, line, size=SIZE_BODY - 0.5, color=COLOR_MUTED)
                set_paragraph_borders(para, edges=("left",), sz=18, color="8FAADC")
            continue

        if kind == "ul":
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.line_spacing = 1.5
            add_inline(para, payload)
            continue

        if kind == "ol":
            num, text = payload
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.75)
            para.paragraph_format.first_line_indent = Cm(-0.75)
            para.paragraph_format.line_spacing = 1.5
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, f"{num}. {text}")
            continue

        if kind == "task":
            done, text = payload
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(0.75)
            para.paragraph_format.first_line_indent = Cm(-0.75)
            para.paragraph_format.line_spacing = 1.4
            para.paragraph_format.space_after = Pt(2)
            add_inline(para, f"{'☑' if done else '☐'} {text}")
            continue

        if kind == "hr":
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(0)
            para.paragraph_format.space_after = Pt(6)
            set_paragraph_borders(para, edges=("bottom",), sz=6)
            continue

        if kind == "code":
            lines = payload or [""]
            # 短代码块整体不跨页，避免逻辑图被拦腰截断
            whole = len(lines) <= 20
            for j, line in enumerate(lines):
                para = doc.add_paragraph()
                pf = para.paragraph_format
                pf.line_spacing = 1.0
                pf.space_before = Pt(6 if j == 0 else 0)
                pf.space_after = Pt(6 if j == len(lines) - 1 else 0)
                if whole:
                    pf.keep_together = True
                    pf.keep_with_next = j < len(lines) - 1
                pf.left_indent = Cm(0.3)
                pf.right_indent = Cm(0.3)
                run = para.add_run(line if line else " ")
                set_run_font(run, FONT_CODE, FONT_CODE, SIZE_CODE)
                shade_paragraph(para, FILL_CODE)
                edges = []
                if j == 0:
                    edges.append("top")
                if j == len(lines) - 1:
                    edges.append("bottom")
                edges += ["left", "right"]
                set_paragraph_borders(para, edges=tuple(edges), sz=4, color="D9D9D9")
            continue

        if kind == "table":
            grid = parse_table(payload)
            if not grid:
                continue
            table = doc.add_table(rows=0, cols=len(grid[0]))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_full_width(table)

            for r, row_cells in enumerate(grid):
                rows = table.add_row()
                cells = rows.cells
                prevent_row_split(rows)
                for c, text in enumerate(row_cells):
                    cell = cells[c]
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.paragraph_format.line_spacing = 1.15
                    para.paragraph_format.space_after = Pt(1)
                    add_inline(para, text, size=SIZE_TABLE,
                               bold=(r == 0))
                    if r == 0:
                        shade_cell(cell, FILL_THEAD)
            repeat_header_row(table.rows[0])
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue


# ---------------------------------------------------------------- 入口
# python-docx 的默认模板在 fontTable.xml / stylesWithEffects.xml 中声明了本机
# 不存在的 Courier 字体（虽是未被使用的潜在样式，但会让渲染器报缺失字体）。
# 保存后统一改写，保持字体表干净。
FONT_TABLE_FIX = {"Courier": FONT_CODE}
FONT_PARTS = ("word/fontTable.xml", "word/stylesWithEffects.xml")


def postprocess_fonts(path):
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path, "r") as zin, \
            zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename in FONT_PARTS:
                text = data.decode("utf-8")
                for old, new in FONT_TABLE_FIX.items():
                    text = text.replace(f'"{old}"', f'"{new}"')
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)


def main():
    here = Path(__file__).resolve().parent
    root = here.parent

    ap = argparse.ArgumentParser(description="Markdown → 规范格式 DOCX")
    ap.add_argument("--input", default=str(root / "项目规划书.md"))
    ap.add_argument("--output", default=str(root / "项目规划书.docx"))
    args = ap.parse_args()

    src, dst = Path(args.input), Path(args.output)
    if not src.is_file():
        sys.exit(f"找不到输入文件：{src}")

    blocks = parse_blocks(src.read_text(encoding="utf-8").splitlines())

    heads = [t for k, t in blocks if k == "h"]
    title = heads[0][1] if heads else src.stem
    subtitle = heads[1][1] if len(heads) > 1 else ""

    doc = build_document()
    set_core_properties(doc, title, subtitle, src.name)
    render(doc, blocks)
    doc.save(str(dst))
    postprocess_fonts(dst)

    kinds = {}
    for k, _ in blocks:
        kinds[k] = kinds.get(k, 0) + 1
    print(f"[OK] {src.name} -> {dst.name}")
    print("     区块统计：" + "，".join(f"{k}={v}" for k, v in sorted(kinds.items())))
    print(f"     段落={len(doc.paragraphs)}，表格={len(doc.tables)}")


if __name__ == "__main__":
    main()
